"""
XLSX Report Exporter - Shopee Dana Cair.
4 Sheets: Ringkasan, Daftar Transaksi, Analisis Margin, Data Modal (Referensi)
"""

import os
import sys
from collections import defaultdict
from datetime import datetime
from typing import Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

from config import HEADER_FILL_COLOR, HEADER_FONT_COLOR, CURRENCY_FORMAT, OUTPUT_DIR


class XLSXExporter:

    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.hdr_font = Font(name="Calibri", size=10, bold=True, color=HEADER_FONT_COLOR)
        self.hdr_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
        self.hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.bold = Font(name="Calibri", size=10, bold=True)
        self.title = Font(name="Calibri", size=14, bold=True)
        self.sub = Font(name="Calibri", size=10, italic=True, color="555555")
        self.border = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

    def _hdr(self, ws, row, count):
        for c in range(1, count + 1):
            cell = ws.cell(row=row, column=c)
            cell.font, cell.fill, cell.alignment, cell.border = self.hdr_font, self.hdr_fill, self.hdr_align, self.border

    def _auto(self, ws):
        for col in ws.columns:
            mx = max((min(len(str(c.value or "")), 40) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = mx + 3

    def _row(self, ws, r, data, money_cols=None, pct_cols=None):
        for c, v in enumerate(data, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = self.border
            if money_cols and c in money_cols:
                cell.number_format = CURRENCY_FORMAT
            if pct_cols and c in pct_cols:
                cell.number_format = '0.0"%"'

    # ── Sheet 1: Ringkasan ───────────────────────────────────────

    def _sheet_ringkasan(self, s, month, year):
        ws = self.wb.create_sheet("Ringkasan")
        ws.cell(row=1, column=1, value="LAPORAN DANA CAIR - SHOPEE").font = self.title
        ws.merge_cells("A1:C1")
        ws.cell(row=2, column=1, value=f"Periode: {_bulan(month)} {year}").font = self.sub
        ws.cell(row=3, column=1, value=f"Dibuat: {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = self.sub

        r = 5
        for c, h in enumerate(["Keterangan", "Jumlah", "Nominal (Rp)"], 1):
            ws.cell(row=r, column=c, value=h)
        self._hdr(ws, r, 3)

        rows = [
            ("Pesanan Selesai (COMPLETED)", s["order_count"], ""),
            ("Total Qty Terjual", f"{s['total_qty_sold']} unit", ""),
            ("", "", ""),
            ("Subtotal Produk (harga jual x qty)", "", s["total_subtotal"]),
            ("(-) Biaya Administrasi", "", -s["total_commission"]),
            ("(-) Biaya Layanan", "", -s["total_service_fee"]),
            ("(-) Biaya Proses Pesanan", "", -s["total_processing_fee"]),
            ("", "", ""),
            ("DANA CAIR (Total Penghasilan)", s["order_count"], s["total_escrow"]),
        ]
        for i, (desc, qty, nom) in enumerate(rows, 1):
            rr = r + i
            ws.cell(row=rr, column=1, value=desc)
            ws.cell(row=rr, column=2, value=qty)
            ws.cell(row=rr, column=3, value=nom)
            if isinstance(nom, (int, float)) and nom != "":
                ws.cell(row=rr, column=3).number_format = CURRENCY_FORMAT
            for c in range(1, 4):
                ws.cell(row=rr, column=c).border = self.border
            if "DANA CAIR" in str(desc):
                for c in range(1, 4):
                    ws.cell(row=rr, column=c).font = self.bold
        self._auto(ws)

    # ── Sheet 2: Daftar Transaksi ────────────────────────────────

    def _sheet_transaksi(self, orders, items):
        ws = self.wb.create_sheet("Daftar Transaksi")

        hdrs = ["No", "Order SN", "Tanggal", "Buyer", "SKU", "Nama Produk", "Qty",
                "Subtotal Produk (Rp)", "Biaya Admin (Rp)", "Biaya Layanan (Rp)",
                "Biaya Proses (Rp)", "Dana Cair (Rp)"]
        for c, h in enumerate(hdrs, 1):
            ws.cell(row=1, column=c, value=h)
        self._hdr(ws, 1, len(hdrs))
        money = {8, 9, 10, 11, 12}

        # Build item lookup by order_sn
        items_by_order = defaultdict(list)
        for it in items:
            items_by_order[it["order_sn"]].append(it)

        row_num = 2
        for i, o in enumerate(orders, 1):
            order_items = items_by_order.get(o["order_sn"], [])
            # SKU + nama: join if multi-item
            if len(order_items) == 1:
                sku = order_items[0]["model_sku"] or order_items[0]["sku"]
                nama = order_items[0]["item_name"]
                qty = order_items[0]["quantity"]
                subtotal = order_items[0]["original_price"]
            elif len(order_items) > 1:
                skus = [it["model_sku"] or it["sku"] for it in order_items]
                sku = ", ".join(s for s in skus if s)
                nama = f"({len(order_items)} produk)"
                qty = sum(it["quantity"] for it in order_items)
                subtotal = sum(it["original_price"] for it in order_items)
            else:
                sku, nama, qty, subtotal = "", "", 0, 0

            self._row(ws, row_num, [
                i, o["order_sn"], o["order_date"], o["buyer_username"],
                sku, nama, qty, subtotal,
                o["commission_fee"], o["service_fee"], o["seller_processing_fee"],
                o["escrow_amount"],
            ], money)
            row_num += 1

        # Total row
        if orders:
            ws.cell(row=row_num, column=7, value="TOTAL").font = self.bold
            for ci, key in [(8, None), (9, "commission_fee"), (10, "service_fee"),
                            (11, "seller_processing_fee"), (12, "escrow_amount")]:
                if key:
                    val = sum(o[key] for o in orders)
                else:
                    val = sum(it["original_price"] for it in items)
                c = ws.cell(row=row_num, column=ci, value=val)
                c.number_format = CURRENCY_FORMAT
                c.font = self.bold
        self._auto(ws)

    # ── Sheet 3: Analisis Margin ─────────────────────────────────

    def _sheet_margin(self, margin_rows, costs):
        ws = self.wb.create_sheet("Analisis Margin")
        ws.cell(row=1, column=1, value="ANALISIS MARGIN PER PRODUK").font = self.title
        ws.merge_cells("A1:O1")
        ws.cell(row=2, column=1, value="Subtotal Produk/unit - Biaya Platform = Dana Cair/unit. Dibandingkan dengan Harga Modal (beli+3%).").font = self.sub
        ws.cell(row=3, column=1, value="Biaya Proses Rp 1.250/order (flat), dibagi rata ke qty. Unit = satuan jual di Shopee (pcs/renceng/pack).").font = self.sub

        hdrs = ["No", "SKU", "Nama Produk", "Varian", "Tier Qty",
                "Jml Order", "Total Qty",
                "Subtotal/unit (Rp)", "Biaya Admin/unit (Rp)", "% Admin",
                "Biaya Layanan/unit (Rp)", "% Layanan", "Biaya Proses/unit (Rp)",
                "Dana Cair/unit (Rp)", "Harga Modal/unit (Rp)", "Selisih (Rp)", "Status"]
        r = 5
        for c, h in enumerate(hdrs, 1):
            ws.cell(row=r, column=c, value=h)
        self._hdr(ws, r, len(hdrs))
        money = {8, 9, 11, 13, 14, 15, 16}

        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(name="Calibri", size=10, bold=True, color="9C0006")
        green_font = Font(name="Calibri", size=10, color="006100")
        sku_header_font = Font(name="Calibri", size=10, bold=True, color="1F4E79")

        # Group margin_rows by SKU
        from collections import OrderedDict
        grouped = OrderedDict()
        for mr in margin_rows:
            key = mr["sku"] or "(tanpa SKU)"
            grouped.setdefault(key, []).append(mr)

        row_num = r + 1
        num = 0
        for sku, rows in grouped.items():
            # Separate SKU groups with 2 blank rows (except first)
            if num > 0:
                row_num += 2

            cost = costs.get(sku)
            modal = cost["modal_per_unit"] if cost else 0

            for mr in rows:
                num += 1
                selisih = mr["avg_cair"] - modal if modal > 0 else 0
                subtotal = mr["avg_listing"]
                pct_admin = (mr["avg_komisi"] / subtotal * 100) if subtotal > 0 else 0
                pct_svc = (mr["avg_svc"] / subtotal * 100) if subtotal > 0 else 0

                if modal == 0:
                    status = "DATA MODAL TIDAK DITEMUKAN"
                elif selisih < 0:
                    status = f"DI BAWAH TARGET (RUGI Rp {abs(selisih):,.0f})"
                else:
                    status = f"DI ATAS TARGET +Rp {selisih:,.0f}"

                self._row(ws, row_num, [
                    num, sku, mr["item_name"], mr["model_name"], mr["tier"],
                    mr["jml_order"], mr["total_qty"],
                    subtotal, mr["avg_komisi"], pct_admin,
                    mr["avg_svc"], pct_svc, mr["avg_proc"],
                    mr["avg_cair"], modal, selisih, status,
                ], money)

                ws.cell(row=row_num, column=10).number_format = '0.0"%"'
                ws.cell(row=row_num, column=12).number_format = '0.0"%"'

                # Highlight RUGI
                if modal > 0 and selisih < 0:
                    for c in range(1, len(hdrs) + 1):
                        ws.cell(row=row_num, column=c).fill = red_fill
                        ws.cell(row=row_num, column=c).font = red_font
                elif modal > 0 and selisih >= 0:
                    ws.cell(row=row_num, column=17).font = green_font

                row_num += 1

        self._auto(ws)

    # ── Sheet 4: Data Modal (Referensi) ──────────────────────────

    def _sheet_modal(self, costs):
        ws = self.wb.create_sheet("Data Modal (Referensi)")
        ws.cell(row=1, column=1, value="HARGA MODAL PER PRODUK").font = self.title
        ws.merge_cells("A1:G1")
        ws.cell(row=2, column=1, value="Sumber: Google Sheet 'ALL PRODUCT'. Hanya produk yang terjual di periode ini.").font = self.sub
        ws.cell(row=3, column=1, value="Rumus: Harga Beli/Karton (in PPN) / Pcs per Karton = Harga Beli/unit. + Margin 3% = Harga Modal/unit.").font = self.sub

        hdrs = ["No", "SKU", "Nama Barang", "Satuan Jual",
                "Pcs/Karton", "Harga Beli/Karton in PPN (Rp)",
                "Harga Beli/unit (Rp)", "Harga Modal/unit +3% (Rp)"]
        r = 5
        for c, h in enumerate(hdrs, 1):
            ws.cell(row=r, column=c, value=h)
        self._hdr(ws, r, len(hdrs))
        money = {6, 7, 8}

        sorted_costs = sorted(costs.items(), key=lambda x: x[0])
        for i, (sku, c) in enumerate(sorted_costs, 1):
            self._row(ws, r + i, [
                i, sku, c["nama_sheet"], c.get("satuan_jual", "pcs"),
                c["pcs_per_karton"], c["beli_per_karton"],
                c["beli_per_unit"], c["modal_per_unit"],
            ], money)

        self._auto(ws)

    # ── Generate ─────────────────────────────────────────────────

    def generate_full_report(self, data, month, year, filename=None):
        print("\n[EXPORT] Generating XLSX...")
        self._sheet_ringkasan(data["summary"], month, year)
        self._sheet_transaksi(data["orders"], data["items"])
        self._sheet_margin(data["margin_rows"], data["costs"])
        self._sheet_modal(data["costs"])

        os.makedirs(OUTPUT_DIR, exist_ok=True)
        fn = filename or f"Laporan_Dana_Cair_Shopee_{_bulan(month)}_{year}.xlsx"
        path = os.path.join(OUTPUT_DIR, fn)
        self.wb.save(path)
        full = os.path.abspath(path)
        print(f"[EXPORT] Saved: {full}")
        return full


def _bulan(m):
    return {1:"Januari",2:"Februari",3:"Maret",4:"April",5:"Mei",6:"Juni",
            7:"Juli",8:"Agustus",9:"September",10:"Oktober",11:"November",12:"Desember"}.get(m, str(m))
